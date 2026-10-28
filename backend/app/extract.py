import base64
import hashlib
import json
import os
import re
import time
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import httpx
from dotenv import load_dotenv
from pydantic import ValidationError
from pypdf import PdfReader

from app.schemas import Item

load_dotenv()

CACHE_DIR = Path(__file__).resolve().parent.parent / "extract_cache"

PROMPT = """You extract deadlines from a university syllabus.
Return ONLY a JSON array, no other text and no markdown.
Each element must have exactly these keys:
  "title": string,
  "type": one of "assignment", "exam", "quiz", "reading",
  "due_date": "YYYY-MM-DD" or null,
  "relative_date": text like "Week 7" or null,
  "grade_weight": number (percent of final grade, 0 if unknown),
  "est_hours": number (your estimate of hours of work),
  "topics": array of up to 6 short strings
Rules:
- Use null when unknown. NEVER guess a date.
- "topics" must come ONLY from the syllabus text, for example the chapters or weekly topics the item covers. For an exam, list the topics from the weeks it covers. If the syllabus doesn't say, use an empty array. NEVER invent topics.
- If something repeats (for example homework due every second Friday), list each occurrence as its own element with its own date and the weight of one occurrence.
- If an item has a week number but no exact date, put the week in relative_date and leave due_date null.
- Include items with no date too (due_date null), such as class participation.

SYLLABUS TEXT:
<<TEXT>>
"""

OCR_PROMPT = (
    "Transcribe all the text in this document exactly as written, including "
    "tables, dates and percentages. Return plain text only, with no commentary."
)

# Used when no API key is set, so you can test the whole pipeline for free.
MOCK_RESPONSE = json.dumps([
    {"title": "Midterm Exam", "type": "exam", "due_date": None,
     "relative_date": "Week 7", "grade_weight": 30, "est_hours": 10,
     "topics": ["Chapter 1", "Chapter 2", "Chapter 3"]},
    {"title": "Essay 1", "type": "assignment", "due_date": "2026-11-20",
     "relative_date": None, "grade_weight": 15, "est_hours": 6, "topics": []},
])


def pdf_to_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def pptx_to_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts).strip()


def file_to_text(filename: str, data: bytes) -> str:
    if filename.lower().endswith(".pptx"):
        return pptx_to_text(data)
    return pdf_to_text(data)


def _post(url, **kwargs):
    """POST with short retries for temporary errors (busy or rate-limited servers)."""
    for attempt in range(3):
        r = httpx.post(url, **kwargs)
        if r.status_code in (429, 500, 502, 503, 504) and attempt < 2:
            time.sleep(2 * (attempt + 1))
            continue
        return r


def ocr_pdf(data: bytes) -> str:
    """Read a scanned PDF with Gemini's built-in document reading."""
    key = os.getenv("LLM_API_KEY", "").strip()
    if not key or os.getenv("LLM_PROVIDER", "anthropic").strip().lower() != "gemini":
        return ""
    fallbacks = [m.strip() for m in os.getenv("LLM_FALLBACK_MODELS", "").split(",") if m.strip()]
    models = [os.getenv("LLM_MODEL", "").strip() or "gemini-flash-latest"] + fallbacks
    body = {
        "contents": [{
            "parts": [
                {"inline_data": {"mime_type": "application/pdf",
                                 "data": base64.b64encode(data).decode()}},
                {"text": OCR_PROMPT},
            ]
        }]
    }
    r = None
    for i, m in enumerate(models):
        r = _post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{m}:generateContent",
            headers={"x-goog-api-key": key},
            json=body,
            timeout=180,
        )
        if r.status_code in (400, 404, 429, 500, 502, 503, 504) and i < len(models) - 1:
            continue
        break
    r.raise_for_status()
    return r.json()["candidates"][0]["content"]["parts"][0]["text"].strip()


def call_llm(prompt: str) -> str:
    key = os.getenv("LLM_API_KEY", "").strip()
    if not key:
        return MOCK_RESPONSE
    provider = os.getenv("LLM_PROVIDER", "anthropic").strip().lower()
    model = os.getenv("LLM_MODEL", "").strip()

    if provider == "anthropic":
        r = _post(
            "https://api.anthropic.com/v1/messages",
            headers={
                "x-api-key": key,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json",
            },
            json={
                "model": model or "claude-sonnet-4-6",
                "max_tokens": 4000,
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=120,
        )
        r.raise_for_status()
        return r.json()["content"][0]["text"]

    if provider == "openai":
        r = _post(
            "https://api.openai.com/v1/chat/completions",
            headers={"Authorization": f"Bearer {key}"},
            json={
                "model": model or "gpt-4o-mini",
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=120,
        )
        r.raise_for_status()
        return r.json()["choices"][0]["message"]["content"]

    if provider == "gemini":
        fallbacks = [m.strip() for m in os.getenv("LLM_FALLBACK_MODELS", "").split(",") if m.strip()]
        models = [model or "gemini-flash-latest"] + fallbacks
        r = None
        for i, m in enumerate(models):
            r = _post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{m}:generateContent",
                headers={"x-goog-api-key": key},
                json={
                    "contents": [{"parts": [{"text": prompt}]}],
                    "generationConfig": {"responseMimeType": "application/json"},
                },
                timeout=120,
            )
            if r.status_code in (400, 404, 429, 500, 502, 503, 504) and i < len(models) - 1:
                continue  # try the next model
            break
        r.raise_for_status()
        return r.json()["candidates"][0]["content"]["parts"][0]["text"]

    raise ValueError(f"Unknown LLM_PROVIDER: {provider}")


def parse_json_array(text: str) -> list[dict]:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    start, end = text.find("["), text.rfind("]")
    if start == -1 or end == -1:
        raise ValueError("No JSON array found")
    rows = json.loads(text[start : end + 1])
    if not isinstance(rows, list):
        raise ValueError("Expected a list")
    return rows


def resolve_relative(item: Item, term_start: date) -> Item:
    """Turn 'Week 7' into a real date and flag it for the student to check."""
    if item.due_date is None and item.relative_date:
        match = re.search(r"week\s*(\d+)", item.relative_date, re.IGNORECASE)
        if match:
            week = int(match.group(1))
            item.due_date = term_start + timedelta(weeks=week - 1)
            item.needs_review = True
    if item.due_date is None:
        item.needs_review = True
    return item


def _load_rows(text: str) -> list[dict]:
    """Return extracted rows, using the cache when this exact text was read before."""
    snippet = text[:12000]
    # "v2|" versions the cache so old results without topics are not reused
    cache_key = hashlib.sha256(("v2|" + snippet).encode()).hexdigest()[:24]
    cache_file = CACHE_DIR / (cache_key + ".json")
    if cache_file.exists():
        try:
            return json.loads(cache_file.read_text(encoding="utf-8"))
        except Exception:
            pass

    prompt = PROMPT.replace("<<TEXT>>", snippet)
    rows = None
    for _attempt in range(2):  # try once, retry once
        try:
            rows = parse_json_array(call_llm(prompt))
            break
        except (ValueError, json.JSONDecodeError):
            continue
    if rows is None:
        raise ValueError("The AI response could not be read as valid JSON")

    if os.getenv("LLM_API_KEY", "").strip():  # never cache the fake mock rows
        try:
            CACHE_DIR.mkdir(exist_ok=True)
            cache_file.write_text(json.dumps(rows), encoding="utf-8")
        except Exception:
            pass
    return rows


def _clean_topics(value) -> list[str]:
    """Keep topics only if they are a list of non-empty strings (max 6)."""
    if not isinstance(value, list):
        return []
    return [str(t).strip() for t in value if str(t).strip()][:6]


def extract_items(text: str, course: str, term_start: date) -> list[Item]:
    rows = _load_rows(text)
    items = []
    for row in rows:
        try:
            if isinstance(row, dict):
                row = {**row, "topics": _clean_topics(row.get("topics"))}
            item = Item(course=course, **row)
        except (ValidationError, TypeError):
            # One bad row must not break the whole sheet
            title = str(row.get("title", "Unknown item")) if isinstance(row, dict) else "Unknown item"
            item = Item(course=course, title=title, type="assignment", needs_review=True)
        items.append(resolve_relative(item, term_start))
    return items